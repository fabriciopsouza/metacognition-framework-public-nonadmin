#!/usr/bin/env python3
"""Canário das planilhas de confirmação e de verificação (B3a-2, ADR-122) — `tools/planilhas_projeto.py`.

Prova, com especificação e perfil SINTÉTICOS (valores literais aqui):
  (a) VERDE — arquivo único completo passa em `verificar`;
  (b) ESTRUTURA — abas e colunas geradas; com os rótulos do modelo real, o cabeçalho sai IGUAL ao modelo
      (critério 3a: mesma estrutura, texto adaptável); 15 de base + 5 por rodada (critério 3b);
  (c) IDA E VOLTA — resposta da área (inclusive com quebra de linha) volta ao arquivo único e está na
      versão seguinte; a tabela de pontos da Parte H ganha as colunas de resposta e o desenho a lê;
  (d) VERSÃO — v1 → v2 apaga a v1 só depois de comparar; resposta não importada reprova e deixa a v1
      intacta; planilha aberta (~$), versão superada, cabeçalho mudado e planilha sem marca são recusados;
  (e) CATEGORIA OBRIGATÓRIA — com impacto regulado, a falta reprova `verificar` e o `gerar` acrescenta;
  (f) SABOTAGENS DE ESTRUTURA — cada uma reprova com o achado dela;
  (g) NOVA RODADA — acrescentar a Rodada 2 depois de respostas importadas não perde nada.

Uso: python tools/test_planilhas_projeto.py   (exit 0 PASS; 1 FAIL; openpyxl ausente = SKIP, exit 0)
"""
import hashlib
import json
import os
import shutil
import sys
import tempfile

sys.dont_write_bytecode = True
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except (AttributeError, ValueError):
    pass

try:
    import openpyxl
except ImportError:
    print("SKIP: openpyxl ausente (tools/requirements-dev.txt)")
    sys.exit(0)

import desenho_processo  # noqa: E402
import planilhas_projeto as pp  # noqa: E402
import regulado  # noqa: E402

PERFIL = {"profile": "teste-sintetico", "testes_obrigatorios_com_impacto": ["integridade de dados", "trilha de auditoria"]}

# Cabeçalhos literais dos modelos reais conferidos em 27/09/2026 (planilha de definições, aba REGRAS; checklist,
# 15 colunas de base). Só a ESTRUTURA é exigida: com estes rótulos o gerado tem de sair idêntico.
MODELO_REGRAS = ["Nº", "Tema", "Regra, como entendemos", "Exemplo", "Como fica no SAP",
                 "Exige alterar o cadastro de muitos materiais?", "Situação hoje no SAP", "É possível? Como? Em massa?",
                 "Links (fonte e tutorial)", "De onde tiramos", "PCP confirma?", "Correção do PCP", "Nome e data"]

V15 = "| Verificação | Onde olhar | Item como aparece na tela | Nome técnico | O que o item faz | Aplica-se a | Regra / valor esperado | Como conferir vários de uma vez | Se estiver errado, o que acontece | Criticidade | Natureza da decisão | Confiança | Valor encontrado | Conclusão da área | Data / Visto |\n" + "|---" * 15 + "|\n"
R6 = "| Verificação | Pergunta de volta | Resposta da área | O que fazer (valor, onde, quem) | Situação | Nome e data |\n|---|---|---|---|---|---|\n"


def ver(n, nome):
    return f"| VER-{n:02d} | tela X | {nome} | CAMPO{n} | faz algo | todos | valor A | consulta | erro | alto | Qualidade | CONFIRMADO | | | |\n"


BASE = ("<!-- spec-unico:v1 -->\n# caso — Especificação\n\n# Parte A — Requisitos\n"
        "- **Impacto regulado:** sim — afeta dado de qualidade\n- **Perfil regulado:** teste-sintetico\n\n- REQ-01: algo\n\n"
        "# Parte H — Processo\n| Passo | Macroação | Descrição | Papel | Gatilho | Saída | Próximo | Regra | Controle |\n"
        "|---|---|---|---|---|---|---|---|---|\n| P-01 | M | faz | Analista | — | — | | — | — |\n\n"
        "### Pontos para aprovação\n| # | Ponto | Por que importa |\n|---|---|---|\n| 1 | aprovar o fluxo | sem isso para |\n\n"
        "# Parte I — Definições a confirmar\n**Para que serve:** confirmar regras\n**Quem responde:** Qualidade\n\n"
        "| Definição | Tema | Regra, como entendemos | Exemplo | Como fica no sistema | Exige alterar muitos registros? | Situação hoje | É possível? Como? Em massa? | Links (fonte e tutorial) | De onde tiramos | Área confirma? | Correção da área | Nome e data |\n"
        + "|---" * 13 + "|\n"
        "| DEF-01 | Demanda | regra um | ex1 | campo A | não | CONFORME | sim | — | reunião | | | |\n"
        "| DEF-02 | Validade | regra dois \\| com barra | ex2 | campo B | sim | DIVERGE | sim, em massa | — | reunião | | | |\n\n"
        "### Valores por categoria\n| Item | Nome técnico | CAT1 | CAT2 | Regra | Área confirma? | Correção da área | Nome e data |\n"
        "|---|---|---|---|---|---|---|---|\n| Tipo de planejamento | TIPO | P2 | PD | DEF-01 | | | |\n\n"
        "### Rótulos\n| Coluna | Rótulo neste projeto |\n|---|---|\n| Como fica no sistema | Como fica no ERPX |\n\n"
        "# Parte J — Roteiro de verificação\n**Para que serve:** conferir o cadastro\n\n### Antes de começar\n- corrigir o cadastro não corrige o passado\n\n"
        "### Bloco: Cadastro básico\n" + V15 + ver(1, "Status") + ver(2, "Grupo") + "\n"
        "### Bloco: integridade de dados\n" + V15 + ver(3, "Log") + "\n"
        "### Bloco: Trilha de Auditoria\n" + V15 + ver(4, "Histórico") + "\n"
        "### Rodada 1 — 23/09/2026\n" + R6 + "| VER-01 | por que branco? | | ajustar | aberta | |\n\n"
        "### Achados do ambiente\n| Achado | O que foi constatado | Por que muda a leitura | O que fazer | Confiança |\n|---|---|---|---|---|\n"
        "| E-01 | grupo sem doc | muda leitura | documentar | INFERIDO |\n\n"
        "### Correções da base\n| Correção | O que constava | Correção aplicada | Motivo |\n|---|---|---|---|\n| C-01 | campo X | campo Y | errado |\n\n"
        "### Controle de versão\n| Versão | Data | Autor | Conteúdo | Base utilizada |\n|---|---|---|---|---|\n| v1.0 | 09/2026 | equipe | base | doc |\n\n"
        "# Decisões\n")

falhas = []


def check(nome, cond, detalhe=""):
    print(("  PASS " if cond else "  FAIL ") + nome + ("" if cond else f" — {detalhe}"))
    if not cond:
        falhas.append(nome)


def troca(velho, novo, texto=BASE):
    assert texto.count(velho) == 1, f"âncora do teste não é única: {velho!r}"
    return texto.replace(velho, novo)


def novo_caso(texto=BASE):
    d = tempfile.mkdtemp()
    spec = os.path.join(d, "spec.md")
    with open(spec, "w", encoding="utf-8") as fh:
        fh.write(texto)
    return d, spec, os.path.join(d, "saida")


def ler(spec):
    with open(spec, encoding="utf-8") as fh:
        return fh.read()


def sha(p):
    with open(p, "rb") as fh:
        return hashlib.sha256(fh.read()).hexdigest()


def cab(xlsx, aba):
    wb = pp._abrir(xlsx)
    return [c.value for c in wb[aba][1] if c.value is not None]


def preencher(xlsx, aba, chave, valores):
    """Simula a área: grava valores {rótulo do cabeçalho: valor} na linha `chave`."""
    wb = pp._abrir(xlsx)
    ws = wb[aba]
    h = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == chave:
            for k, v in valores.items():
                row[h.index(k)].value = v
    wb.save(xlsx)


def silencioso(*_a, **_k):
    pass


def recusa(fn):
    """Mensagem da Recusa levantada por fn (vazio se não recusou)."""
    try:
        fn()
    except pp.Recusa as e:
        return str(e)
    return ""


def ultima(out, tipo):
    return os.path.join(out, f"{tipo}_v{pp._ultima(out, tipo)}.xlsx")


def celula(xlsx, aba, chave, rotulo):
    wb = pp._abrir(xlsx)
    ws = wb[aba]
    h = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == chave:
            return row[h.index(rotulo)]
    return None


def main():
    pdir = tempfile.mkdtemp()
    with open(os.path.join(pdir, "compliance-profile-teste-sintetico.json"), "w", encoding="utf-8") as fh:
        json.dump(PERFIL, fh)
    original = regulado.PERFIS_DIR
    regulado.PERFIS_DIR = pdir
    tmp = []
    try:
        print("(a) verde")
        d, spec, out = novo_caso(); tmp.append(d)
        _, ach = pp.verificar(spec)
        check("arquivo único completo passa", ach == [], ach)

        print("(b) estrutura")
        check("gerar sai com 0", pp.main(["gerar", spec, "--out-dir", out]) == 0)
        c1, v1 = os.path.join(out, "confirmacao_v1.xlsx"), os.path.join(out, "verificacao_v1.xlsx")
        check("as duas planilhas v1 existem", os.path.isfile(c1) and os.path.isfile(v1), os.listdir(out))
        wb = pp._abrir(c1)
        check("abas da confirmação", wb.sheetnames == ["LEIA", "REGRAS", "VALORES POR CATEGORIA"], wb.sheetnames)
        h = cab(c1, "REGRAS")
        check("REGRAS com 13 colunas e o rótulo do projeto", len(h) == 13 and h[4] == "Como fica no ERPX", h)
        ids = [r[0].value for r in wb["REGRAS"].iter_rows(min_row=2)]
        check("linhas: DEF-01, DEF-02 e o ponto H-1 da Parte H", ids == ["DEF-01", "DEF-02", "H-1"], ids)
        check("barra escapada volta como '|'", wb["REGRAS"]["C3"].value == "regra dois | com barra", wb["REGRAS"]["C3"].value)
        wb = pp._abrir(v1)
        check("abas da verificação", wb.sheetnames == ["COMO USAR", "ANTES DE COMEÇAR", "CHECKLIST", "ACHADOS DO AMBIENTE",
                                                         "CORREÇÕES DA BASE", "CONTROLE DE VERSÃO"], wb.sheetnames)
        hv = cab(v1, "CHECKLIST")
        check("CHECKLIST: 15 de base + 5 da rodada 1", len(hv) == 20 and hv[15].startswith("Rodada 1 (23/09/2026)"), hv)
        como = [r[0].value for r in wb["COMO USAR"].iter_rows()]
        check("COMO USAR traz a instrução de acrescentar rodada", "Como acrescentar uma rodada" in como, como)
        # critério 3a: com os rótulos do modelo real, o cabeçalho gerado é o do modelo, posição a posição
        rot = "".join(f"| {n} | {m} |\n" for n, m in zip(pp.COLS_DEF, MODELO_REGRAS) if n != m)
        d2, spec2, out2 = novo_caso(troca("| Como fica no sistema | Como fica no ERPX |\n", rot)); tmp.append(d2)
        pp.main(["gerar", spec2, "--out-dir", out2, "--so", "confirmacao"])
        h2 = cab(os.path.join(out2, "confirmacao_v1.xlsx"), "REGRAS")
        check("rótulos do modelo real reproduzem o cabeçalho real", h2 == MODELO_REGRAS, h2)

        print("(c) ida e volta")
        preencher(c1, "REGRAS", "DEF-01", {"Área confirma?": "Com correção", "Correção da área": "linha 1\nlinha 2",
                                           "Nome e data": "Ana, 27/09/2026"})
        preencher(c1, "REGRAS", "H-1", {"Área confirma?": "Sim", "Nome e data": "Rui, 27/09/2026"})
        preencher(c1, "VALORES POR CATEGORIA", "Tipo de planejamento", {"Área confirma?": "Não", "Correção da área": "PD também"})
        esperado = pp.respostas(c1, "confirmacao", pp.carregar(spec))
        pp.importar(spec, c1, out, log=silencioso)
        t = ler(spec)
        check("resposta com quebra de linha gravada como <br>", "| Com correção | linha 1<br>linha 2 | Ana, 27/09/2026 |" in t)
        check("barra de outra célula continua escapada", "regra dois \\| com barra" in t)
        check("pontos da Parte H ganharam as colunas de resposta",
              "| # | Ponto | Por que importa | Área confirma? | Correção da área | Nome e data |" in t and
              "| 1 | aprovar o fluxo | sem isso para | Sim |  | Rui, 27/09/2026 |" in t)
        _, pontos, _, ach_h = desenho_processo.carregar(spec)
        check("desenho do processo lê os pontos ampliados", len(pontos) == 1 and not ach_h, (pontos, ach_h))
        c2 = os.path.join(out, "confirmacao_v2.xlsx")
        check("v2 criada, v1 apagada, _obsoleto removida", os.path.isfile(c2) and not os.path.exists(c1)
              and not os.path.exists(os.path.join(out, "_obsoleto")), os.listdir(out))
        got = pp.respostas(c2, "confirmacao", pp.carregar(spec))
        check("toda resposta da v1 está na v2, idêntica", got == esperado, (got, esperado))
        preencher(v1, "CHECKLIST", "VER-01", {"Valor encontrado": "branco", "Rodada 1 (23/09/2026) · Resposta da área": "é o padrão"})
        preencher(v1, "CHECKLIST", "VER-03", {"Rodada 1 (23/09/2026) · Resposta da área": "sem log",
                                              "Rodada 1 (23/09/2026) · Nome e data": "Ana, 27/09/2026"})
        pp.importar(spec, v1, out, log=silencioso)
        t = ler(spec)
        check("verificação de base gravada", "| CONFIRMADO | branco |  |  |" in t)
        check("rodada: linha existente atualizada", "| VER-01 | por que branco? | é o padrão | ajustar | aberta |  |" in t)
        check("rodada: linha nova acrescentada na rodada", "| VER-03 |  | sem log |  |  | Ana, 27/09/2026 |" in t)

        print("(d) versão")
        # arquivo aberto testado ANTES de haver resposta pendente: a única recusa possível é a do ~$
        c2h = sha(c2)
        trava = os.path.join(out, "~$confirmacao_v2.xlsx")
        open(trava, "w").close()
        msg = recusa(lambda: pp.gerar_tipo(spec, "confirmacao", out, log=silencioso))
        check("planilha aberta (~$) é recusada pelo motivo certo", "aberta no Excel" in msg and sha(c2) == c2h
              and not os.path.exists(os.path.join(out, "confirmacao_v3.xlsx")), msg)
        msg = recusa(lambda: pp.importar(spec, c2, out, log=silencioso))
        check("importar planilha aberta é recusado pelo motivo certo", "aberta no Excel" in msg, msg)
        os.remove(trava)
        preencher(c2, "REGRAS", "DEF-02", {"Área confirma?": "Sim"})
        c2h = sha(c2)
        msg = recusa(lambda: pp.gerar_tipo(spec, "confirmacao", out, log=silencioso))
        check("resposta não importada reprova o gerar", "ainda não gravadas" in msg and "DEF-02" in msg, msg)
        check("v2 intacta no lugar, sem v3, sem _obsoleto", os.path.isfile(c2) and sha(c2) == c2h
              and not os.path.exists(os.path.join(out, "confirmacao_v3.xlsx"))
              and not os.path.exists(os.path.join(out, "_obsoleto")), os.listdir(out))
        antes = ler(spec)
        pp.importar(spec, c2, out, log=silencioso)
        velha = os.path.join(d, "v2-devolvida.xlsx")
        shutil.copy(os.path.join(out, "confirmacao_v3.xlsx"), velha)
        pp.main(["gerar", spec, "--out-dir", out, "--so", "confirmacao"])
        rc = pp.main(["importar", spec, "--planilha", velha, "--out-dir", out])
        check("versão superada (v3 com v4 vigente) é recusada", rc == 1)
        c4 = os.path.join(out, "confirmacao_v4.xlsx")
        wb = pp._abrir(c4)
        wb["REGRAS"]["B1"] = "Assunto"
        wb.save(c4)
        antes = ler(spec)
        rc = pp.main(["importar", spec, "--planilha", c4, "--out-dir", out])
        check("cabeçalho mudado pela área é recusado e o arquivo único não muda", rc == 1 and ler(spec) == antes)
        crua = os.path.join(d, "crua.xlsx")
        openpyxl.Workbook().save(crua)
        check("planilha sem a marca de versão é recusada",
              pp.main(["importar", spec, "--planilha", crua, "--out-dir", out]) == 1)

        print("(e) categoria obrigatória")
        sem = troca("### Bloco: Trilha de Auditoria\n" + V15 + ver(4, "Histórico") + "\n", "")
        d3, spec3, out3 = novo_caso(sem); tmp.append(d3)
        _, ach = pp.verificar(spec3)
        check("falta da categoria reprova verificar", any("trilha de auditoria" in a for a in ach), ach)
        rc = pp.main(["gerar", spec3, "--out-dir", out3])
        t3 = ler(spec3)
        check("gerar acrescenta o bloco e passa", rc == 0 and "### Bloco: trilha de auditoria" in t3
              and "| VER-04 |" in t3, rc)
        check("bloco entra antes da rodada", t3.index("### Bloco: trilha de auditoria") < t3.index("### Rodada 1"))
        d4, spec4, _ = novo_caso(troca("- **Impacto regulado:** sim — afeta dado de qualidade", "- **Impacto regulado:** não", sem))
        tmp.append(d4)
        _, ach = pp.verificar(spec4)
        check("sem impacto regulado, nenhuma categoria é exigida", ach == [], ach)

        print("(f) sabotagens de estrutura")
        casos = [
            ("definição com 12 colunas", troca("| reunião | | | |\n| DEF-02", "| reunião | | |\n| DEF-02"), "DEF-01 com 12 colunas"),
            ("cabeçalho da Parte I alterado", troca("| Definição | Tema |", "| Definição | Assunto |"), "cabeçalho difere"),
            ("rodada cita verificação inexistente", troca("| VER-01 | por que branco?", "| VER-99 | por que branco?"), "VER-99 não existe"),
            ("rodada sem data", troca("### Rodada 1 — 23/09/2026", "### Rodada 1"), "Rodada N — dd/mm/aaaa"),
            ("rodada fora de sequência", troca("### Rodada 1 — 23/09/2026", "### Rodada 2 — 23/09/2026"), "fora de sequência"),
            ("verificação repetida", troca(ver(4, "Histórico"), ver(3, "Histórico")), "VER-03 repetida"),
            ("rótulo de coluna inexistente", troca("| Como fica no sistema | Como fica no ERPX |", "| Coluna X | Y |"), "Coluna X"),
            ("valores por categoria sem colunas de resposta", troca("| Regra | Área confirma? | Correção da área | Nome e data |\n|---|---|---|---|---|---|---|---|",
                                                                     "| Regra |\n|---|---|---|---|---|"), "Valores por categoria"),
        ]
        for nome, texto, trecho in casos:
            dd, s, _ = novo_caso(texto); tmp.append(dd)
            _, ach = pp.verificar(s)
            check(nome, any(trecho in a for a in ach), ach)

        print("(g) nova rodada depois de respostas importadas")
        t = ler(spec)
        t = t.replace("### Achados do ambiente", "### Rodada 2 — 28/09/2026\n" + R6 + "| VER-02 | e o grupo? | | | aberta | |\n\n### Achados do ambiente")
        with open(spec, "w", encoding="utf-8") as fh:
            fh.write(t)
        rc = pp.main(["gerar", spec, "--out-dir", out, "--so", "verificacao"])
        vn = os.path.join(out, f"verificacao_v{pp._ultima(out, 'verificacao')}.xlsx")
        hv = cab(vn, "CHECKLIST")
        check("rodada 2 vira 5 colunas novas sem perder a rodada 1", rc == 0 and len(hv) == 25
              and hv[20].startswith("Rodada 2 (28/09/2026)"), (rc, hv[15:]))
        wb = pp._abrir(vn)
        linha = next(r for r in wb["CHECKLIST"].iter_rows(min_row=2, values_only=True) if r[0] == "VER-03")
        check("resposta da rodada 1 continua na versão nova", "sem log" in linha, linha)

        print("(h) achados da revisão de 27/09 — cada um fechado com caso próprio")
        d5, spec5, out5 = novo_caso(); tmp.append(d5)
        pp.gerar_tipo(spec5, "confirmacao", out5, log=silencioso)
        c = ultima(out5, "confirmacao")
        cel = celula(c, "REGRAS", "DEF-01", "Nome e data")
        check("coluna de resposta sai em formato texto (@)", cel.number_format == "@", cel.number_format)
        preencher(c, "REGRAS", "DEF-01", {"Área confirma?": "Sim", "Correção da área": "=03/2026",
                                          "Nome e data": "Ana, 28/09/2026"})
        preencher(c, "REGRAS", "DEF-02", {"Área confirma?": "Com correção", "Correção da área": "usar <br> literal",
                                          "Nome e data": "Rui, 28/09/2026"})
        pp.importar(spec5, c, out5, log=silencioso)
        t5 = ler(spec5)
        check("resposta que começa com '=' volta como texto", "| Sim | =03/2026 | Ana, 28/09/2026 |" in t5)
        check("'<br>' digitado vira &lt;br&gt; no arquivo único", "usar &lt;br&gt; literal" in t5)
        c = ultima(out5, "confirmacao")
        cf = celula(c, "REGRAS", "DEF-01", "Correção da área")
        check("na versão seguinte o '=' continua texto, não fórmula", cf.value == "=03/2026" and cf.data_type == "s",
              (cf.value, cf.data_type))
        check("'<br>' literal volta idêntico na versão seguinte",
              celula(c, "REGRAS", "DEF-02", "Correção da área").value == "usar <br> literal")
        d11, spec11, out11 = novo_caso(""); tmp.append(d11)
        check("arquivo único vazio não gera nada, mesmo chamado direto",
              "achados de estrutura" in recusa(lambda: pp.gerar_tipo(spec11, "confirmacao", out11, log=silencioso))
              and not os.path.isdir(out11))
        check("gerar de novo passa (sem recusa perpétua)", recusa(lambda: pp.gerar_tipo(spec5, "confirmacao", out5,
                                                                                          log=silencioso)) == "")
        # rótulo alterado no arquivo único depois de resposta importada: nada se perde, nada é recusado
        novo_rotulo = ler(spec5).replace("| Como fica no sistema | Como fica no ERPX |", "| Como fica no sistema | Onde fica |")
        with open(spec5, "w", encoding="utf-8") as fh:  # ler ANTES de abrir para escrita (o "w" zera o arquivo)
            fh.write(novo_rotulo)
        msg = recusa(lambda: pp.gerar_tipo(spec5, "confirmacao", out5, log=silencioso))
        c = ultima(out5, "confirmacao")
        check("rótulo mudado entre versões: gera e preserva as respostas", msg == ""
              and celula(c, "REGRAS", "DEF-01", "Correção da área").value == "=03/2026", msg)
        # linha removida com resposta: recusa com causa própria; --descartar sem data não vale; com data guarda a anterior
        sem01 = "\n".join(ln for ln in ler(spec5).splitlines() if not ln.startswith("| DEF-01 |")) + "\n"
        with open(spec5, "w", encoding="utf-8") as fh:
            fh.write(sem01)
        antes_v = pp._ultima(out5, "confirmacao")
        msg = recusa(lambda: pp.gerar_tipo(spec5, "confirmacao", out5, log=silencioso))
        check("linha removida com resposta: recusa própria, com a saída", "saíram do arquivo único" in msg
              and "--descartar" in msg and pp._ultima(out5, "confirmacao") == antes_v, msg)
        msg = recusa(lambda: pp.gerar_tipo(spec5, "confirmacao", out5, log=silencioso, descartar="porque sim"))
        check("--descartar sem nome e data é recusado", "data válida" in msg, msg)
        msg = recusa(lambda: pp.gerar_tipo(spec5, "confirmacao", out5, log=silencioso,
                                           descartar="Fabricio: DEF-01 saiu do escopo, 28/09/2026"))
        guardada = os.path.join(out5, "_obsoleto", f"confirmacao_v{antes_v}.xlsx")
        check("--descartar com data gera e GUARDA a anterior em _obsoleto", msg == "" and os.path.isfile(guardada)
              and pp._ultima(out5, "confirmacao") == antes_v + 1, (msg, os.listdir(out5)))
        # falha ao apagar a anterior: a nova vale, a anterior fica, o comando avisa e não explode
        d6, spec6, out6 = novo_caso(); tmp.append(d6)
        pp.gerar_tipo(spec6, "confirmacao", out6, log=silencioso)
        avisos = []
        real = pp.os.remove

        def remove_trava(p, real=real):
            if "_obsoleto" in p:
                raise PermissionError("em uso")
            real(p)
        pp.os.remove = remove_trava
        try:
            msg = recusa(lambda: pp.gerar_tipo(spec6, "confirmacao", out6, log=avisos.append))
        finally:
            pp.os.remove = real
        check("falha ao apagar a anterior: v2 vale, v1 fica em _obsoleto e há AVISO", msg == ""
              and os.path.isfile(os.path.join(out6, "confirmacao_v2.xlsx"))
              and os.path.isfile(os.path.join(out6, "_obsoleto", "confirmacao_v1.xlsx"))
              and any("AVISO" in a for a in avisos), (msg, avisos))
        # CRLF preservado pelo importar e pelo acrescentar_categorias
        crlf = BASE.replace("\n", "\r\n")
        d7, spec7, out7 = novo_caso(); tmp.append(d7)
        with open(spec7, "w", encoding="utf-8", newline="") as fh:
            fh.write(crlf.replace("### Bloco: Trilha de Auditoria\r\n", "### Bloco: outro nome\r\n"))
        pp.acrescentar_categorias(spec7)
        pp.gerar_tipo(spec7, "confirmacao", out7, log=silencioso)
        c = ultima(out7, "confirmacao")
        preencher(c, "REGRAS", "DEF-01", {"Área confirma?": "Sim", "Nome e data": "Ana, 28/09/2026"})
        pp.importar(spec7, c, out7, log=silencioso)
        with open(spec7, "rb") as fh:
            b = fh.read()
        check("arquivo em CRLF continua 100% CRLF depois de acrescentar e importar",
              b.count(b"\r\n") == b.count(b"\n") and b"### Bloco: trilha de auditoria" in b and b"Ana, 28/09/2026" in b,
              (b.count(b"\r\n"), b.count(b"\n")))

        print("(i) comandos determinísticos: situacao e rodada")
        rel, pend = pp.situacao(spec7)
        check("situacao conta a confirmação assinada", "confirmação: 1 de 4 linha(s) com resposta assinada" in rel, rel)
        check("situacao lista pendência da definição sem resposta", any("DEF-02" in p for p in pend), pend)
        check("situacao --exigir reprova com pendência", pp.main(["situacao", spec7, "--exigir"]) == 1)
        check("situacao sem --exigir só informa", pp.main(["situacao", spec7]) == 0)
        tudo = ler(spec7)
        tudo = tudo.replace("| reunião | | | |", "| reunião | Sim | | Ana, 28/09/2026 |")  # DEF-02
        tudo = tudo.replace("| DEF-01 | | | |", "| DEF-01 | Sim | | Ana, 28/09/2026 |")  # linha de valores
        tudo = tudo.replace("| 1 | aprovar o fluxo | sem isso para |", "| 1 | aprovar o fluxo | sem isso para | Sim |  | Rui, 28/09/2026 |")
        d8, spec8, _ = novo_caso(tudo); tmp.append(d8)
        rel8, pend8 = pp.situacao(spec8)
        conf8 = [p for p in pend8 if p.startswith(("definição", "ponto", "valor"))]
        check("confirmação toda assinada com 'Sim' não tem pendência", conf8 == [], conf8)
        d9, spec9, _ = novo_caso(troca("| DEF-01 | Demanda | regra um | ex1 | campo A | não | CONFORME | sim | — | reunião | | | |",
                                       "| DEF-01 | Demanda | regra um | ex1 | campo A | não | CONFORME | sim | — | reunião | Não | outra regra | Ana, 28/09/2026 |"))
        tmp.append(d9)
        _, pend9 = pp.situacao(spec9)
        check("resposta 'Não' assinada vira pendência a tratar", any("DEF-01" in p and "tratar" in p for p in pend9), pend9)
        linha01 = "| DEF-01 | Demanda | regra um | ex1 | campo A | não | CONFORME | sim | — | reunião | | | |"
        for nome, assinatura in (("nome sem data", "Ana"), ("data sem nome", "27/09/2026"),
                                 ("data impossível", "Ana, 31/02/2026")):
            dn, sn, _ = novo_caso(troca(linha01, linha01[:-len("| | | |")] + f"| Sim | | {assinatura} |")); tmp.append(dn)
            _, pn = pp.situacao(sn)
            check(f"'Sim' com {nome} continua pendente", any("DEF-01" in p and "sem resposta assinada" in p for p in pn), pn)
        d10, spec10, _ = novo_caso(); tmp.append(d10)
        n = pp.nova_rodada(spec10, "28/09/2026", ["VER-02", "VER-04", "VER-02"], log=silencioso)
        check("--ids repetido entra uma vez", ler(spec10).count("| VER-02 |  |  |  |  |  |") == 1)
        t10 = ler(spec10)
        _, ach10 = pp.verificar(spec10)
        check("rodada acrescenta a próxima (2) com as verificações pedidas", n == 2 and "### Rodada 2 — 28/09/2026" in t10
              and "| VER-02 |  |  |  |  |  |" in t10 and ach10 == [], (n, ach10))
        check("rodada entra depois da rodada 1 e antes dos achados",
              t10.index("### Rodada 1") < t10.index("### Rodada 2") < t10.index("### Achados do ambiente"))
        check("rodada com verificação inexistente é recusada",
              "VER-99" in recusa(lambda: pp.nova_rodada(spec10, "28/09/2026", ["VER-99"], log=silencioso)))
        check("rodada com data inválida é recusada",
              "dd/mm/aaaa" in recusa(lambda: pp.nova_rodada(spec10, "31/02/2026", ["VER-01"], log=silencioso)))
    finally:
        regulado.PERFIS_DIR = original
        for p in tmp + [pdir]:
            shutil.rmtree(p, ignore_errors=True)
    print(f"\nRESULTADO: {'FAIL' if falhas else 'PASS'} ({len(falhas)} falha(s))")
    return 1 if falhas else 0


if __name__ == "__main__":
    sys.exit(main())
